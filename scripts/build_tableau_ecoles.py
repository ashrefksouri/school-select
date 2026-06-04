#!/usr/bin/env python3
"""Generate data/tableau-ecoles.xlsx from public/data/schools.json.

Adds two computed columns at the front:
  - template_version: "v4" if public/{slug}/index.html has > 1000 lines, else "v3"
  - tier: "verified" if slug is in the verified list, else "basic"
"""

import json
import os
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SCHOOLS_JSON = ROOT / "public" / "data" / "schools.json"
PUBLIC_DIR = ROOT / "public"
OUTPUT = ROOT / "data" / "tableau-ecoles.xlsx"

VERIFIED_SLUGS = [
    "british-international-school-of-tunis",
    "oxford-international-school-of-tunis",
    "acst",
    "ecole-internationale-de-carthage-eic",
    "ecole-internationale-de-tunis-eit-2",
    "british-academy-of-tunis",
    "college-international-francais-jean-racine",
    "gsing",
    "lycee-francais-international-de-sousse-mhamed-driss",
    "lycee-louis-pasteur",
    "lycee-louis-pasteur-gammarth",
    "groupe-scolaire-rene-descartes-les-berges-du-lac",
    "cis-international-school-of-tunis",
    "international-maarif-schools-of-tunisia",
    "college-lycee-international-francais-jules-verne",
]

# Alias: schools.json slug -> folder slug under public/
# These rows have a schools.json slug that diverges from the actual public/{folder}/index.html.
SLUG_ALIASES = {
    "american-cooperative-school-of-tunis-acst": "acst",
    "advanced-for-education-british-academy-of-tunis": "british-academy-of-tunis",
}


def count_lines(path: Path) -> int | None:
    if not path.exists():
        return None
    with path.open("rb") as f:
        return sum(1 for _ in f)


def stringify(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (list, dict)):
        try:
            return json.dumps(value, ensure_ascii=False)
        except (TypeError, ValueError):
            return str(value)
    return value


def main() -> int:
    with SCHOOLS_JSON.open(encoding="utf-8") as f:
        schools = json.load(f)

    base_keys = list(schools[0].keys())
    for s in schools:
        for k in s.keys():
            if k not in base_keys:
                base_keys.append(k)

    headers = ["template_version", "tier", *base_keys]

    verified_set = set(VERIFIED_SLUGS)
    seen_slugs = set()
    rows = []
    v4_count = 0
    v3_count = 0
    verified_count = 0
    basic_count = 0

    for school in schools:
        slug = school.get("slug", "")
        seen_slugs.add(slug)
        folder_slug = SLUG_ALIASES.get(slug, slug)
        seen_slugs.add(folder_slug)

        lines = count_lines(PUBLIC_DIR / slug / "index.html")
        if lines is None and folder_slug != slug:
            lines = count_lines(PUBLIC_DIR / folder_slug / "index.html")
        template_version = "v4" if lines is not None and lines > 1000 else "v3"
        if template_version == "v4":
            v4_count += 1
        else:
            v3_count += 1

        tier = "verified" if (slug in verified_set or folder_slug in verified_set) else "basic"
        if tier == "verified":
            verified_count += 1
        else:
            basic_count += 1

        row = [template_version, tier]
        for k in base_keys:
            row.append(stringify(school.get(k)))
        rows.append(row)

    wb = Workbook()
    ws = wb.active
    ws.title = "ecoles"

    header_font = Font(bold=True, color="FFFFFFFF")
    header_fill = PatternFill("solid", fgColor="FF1F4E78")
    verified_fill = PatternFill("solid", fgColor="FFE2F0D9")
    v4_font = Font(bold=True)

    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append(row)

    for row_idx in range(2, len(rows) + 2):
        template_version = ws.cell(row=row_idx, column=1).value
        tier = ws.cell(row=row_idx, column=2).value
        if tier == "verified":
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = verified_fill
        if template_version == "v4":
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                existing = cell.font
                cell.font = Font(
                    name=existing.name,
                    size=existing.size,
                    bold=True,
                    color=existing.color,
                )

    last_col_letter = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col_letter}{len(rows) + 1}"
    ws.freeze_panes = "A2"

    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(header))
        for row_idx in range(2, len(rows) + 2):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val is None:
                continue
            s = str(val)
            if len(s) > max_len:
                max_len = len(s)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 60)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)

    missing_verified = [s for s in VERIFIED_SLUGS if s not in seen_slugs]

    print("=== Rapport ===")
    print(f"Total écoles : {len(schools)}")
    print(f"  verified : {verified_count}")
    print(f"  basic    : {basic_count}")
    print(f"  v4       : {v4_count}")
    print(f"  v3       : {v3_count}")
    if missing_verified:
        print(f"Slugs verified introuvables dans schools.json ({len(missing_verified)}) :")
        for s in missing_verified:
            print(f"  - {s}")
    else:
        print("Tous les slugs verified présents dans schools.json.")
    print(f"Fichier généré : {OUTPUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
